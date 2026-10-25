#!/usr/bin/env python
"""Reconstruct an approximate PPTX template from the SK hynix PDF deck (no original master available).

Usage: python pdf-to-pptx.py <deck.pdf> <outDir>
Writes <outDir>/skhynix-reconstructed.pptx plus the cropped PNGs it embeds.
Needs pdfplumber (with Pillow) and python-pptx.

Everything here is measured from one PDF export: colours are sampled pixels, positions are
text boxes from pdfplumber, the halftone art and the logo are raster crops. It is an
approximation of the original slide master, not the master itself.
"""
import os
import sys
from collections import Counter

import pdfplumber
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.shapes.picture import CT_Picture

DPI = 200
S = DPI / 72  # PDF points -> render pixels
EMU = 12700  # EMU per point
FONT_MAJOR = '나눔스퀘어 네오'
FONT_MINOR = '나눔스퀘어'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def hexof(rgb):
    return '%02X%02X%02X' % tuple(rgb[:3])


def dominant(im, box, bg, min_dist=60):
    """Most common colour in a PDF-point box, ignoring pixels close to the background."""
    x0, y0, x1, y1 = [int(v * S) for v in box]
    c = Counter()
    for y in range(y0, y1):
        for x in range(x0, x1):
            p = im.getpixel((x, y))
            if sum(abs(a - b) for a, b in zip(p, bg)) >= min_dist:
                c[p] += 1
    return c.most_common(1)[0][0]


def keyed_crop(im, box, bg, blank=()):
    """Crop a PDF-point box; pixels near the background colour become transparent."""
    x0, y0, x1, y1 = [int(v * S) for v in box]
    out = im.crop((x0, y0, x1, y1)).convert('RGBA')
    px = out.load()
    blanks = [[int(v * S) - o for v, o in zip(b, (x0, y0, x0, y0))] for b in blank]
    for y in range(out.height):
        for x in range(out.width):
            if any(bx0 <= x < bx1 and by0 <= y < by1 for bx0, by0, bx1, by1 in blanks):
                px[x, y] = (0, 0, 0, 0)
                continue
            r, g, b, _ = px[x, y]
            d = abs(r - bg[0]) + abs(g - bg[1]) + abs(b - bg[2])
            px[x, y] = (r, g, b, max(0, min(255, (d - 10) * 6)))
    return out


def emu(*pts):
    return [int(round(v * EMU)) for v in pts]


def set_xfrm(sp, x, y, w, h):
    xfrm = sp.find('.//' + qn('a:xfrm'))
    if xfrm is None:
        sppr = sp.find(qn('p:spPr'))
        xfrm = etree.SubElement(sppr, qn('a:xfrm'))
        sppr.insert(0, xfrm)
    for c in list(xfrm):
        xfrm.remove(c)
    X, Y, W, H = emu(x, y, w, h)
    etree.SubElement(xfrm, qn('a:off'), x=str(X), y=str(Y))
    etree.SubElement(xfrm, qn('a:ext'), cx=str(W), cy=str(H))


def set_lst_style(sp, size=None, bold=None, scheme=None, algn=None):
    """Level-1 text style on a layout placeholder (what the extractor reads first)."""
    tx = sp.find(qn('p:txBody'))
    lst = tx.find(qn('a:lstStyle'))
    for c in list(lst):
        lst.remove(c)
    lvl = etree.SubElement(lst, qn('a:lvl1pPr'))
    if algn:
        lvl.set('algn', algn)
    d = etree.SubElement(lvl, qn('a:defRPr'))
    if size:
        d.set('sz', str(int(size * 100)))
    if bold is not None:
        d.set('b', '1' if bold else '0')
    if scheme:
        f = etree.SubElement(d, qn('a:solidFill'))
        etree.SubElement(f, qn('a:schemeClr'), val=scheme)


def scale_x(part_el, f):
    for xfrm in part_el.iter(qn('a:xfrm')):
        off, ext = xfrm.find(qn('a:off')), xfrm.find(qn('a:ext'))
        if off is not None:
            off.set('x', str(int(int(off.get('x')) * f)))
        if ext is not None:
            ext.set('cx', str(int(int(ext.get('cx')) * f)))


def ph_type(el):
    return el.find('.//' + qn('p:ph')).get('type', 'obj')


def placeholder(layout, *types):
    return next((sp._element for sp in layout.placeholders if ph_type(sp._element) in types), None)


def drop_placeholders(layout, keep):
    for sp in list(layout.placeholders):
        if ph_type(sp._element) not in keep:
            sp._element.getparent().remove(sp._element)


def add_picture(layout, path, x, y, w, h, name):
    _, rid = layout.part.get_or_add_image_part(path)
    tree = layout.shapes._spTree
    ids = [int(e.get('id')) for e in tree.iter(qn('p:cNvPr'))]
    pic = CT_Picture.new_pic(max(ids + [1]) + 1, name, '', rid, *emu(x, y, w, h))
    tree.insert(2, pic)  # behind the placeholders (after nvGrpSpPr, grpSpPr)


def add_rect(layout, x, y, w, h, scheme, name):
    tree = layout.shapes._spTree
    ids = [int(e.get('id')) for e in tree.iter(qn('p:cNvPr'))]
    X, Y, W, H = emu(x, y, w, h)
    xml = (
        f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}">'
        f'<p:nvSpPr><p:cNvPr id="{max(ids + [1]) + 1}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{X}" y="{Y}"/><a:ext cx="{W}" cy="{H}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="{scheme}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
        f'</p:sp>'
    )
    tree.insert(2, etree.fromstring(xml))


def set_bg(el, scheme):
    csld = el.find(qn('p:cSld'))
    old = csld.find(qn('p:bg'))
    if old is not None:
        csld.remove(old)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}">'
        f'<p:bgPr><a:solidFill><a:schemeClr val="{scheme}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'
    )
    csld.insert(0, bg)


def main():
    if len(sys.argv) != 3:
        print('usage: pdf-to-pptx.py <deck.pdf> <outDir>', file=sys.stderr)
        sys.exit(2)
    pdf_path, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    pdf = pdfplumber.open(pdf_path)
    render = lambda i: pdf.pages[i].to_image(resolution=DPI).original.convert('RGB')
    p1, p2, p4, p8 = render(0), render(1), render(3), render(7)

    # ---- colours (pixel samples)
    bg = p1.getpixel((20, 20))                          # cover off-white
    red = dominant(p1, (51, 368, 575, 397), bg)         # cover title
    dark = dominant(p2, (58, 366, 588, 378), bg)        # disclaimer body text
    orange = dominant(p1, (400, 40, 520, 220), bg)      # halftone orange dots
    pink = dominant(p1, (560, 20, 760, 220), bg)        # halftone red dots
    grey = dominant(p8, (505, 155, 685, 158), (255, 255, 255))  # grey label bar
    salmon = dominant(p4, (445, 152, 515, 166), (255, 255, 255))  # "CHAPTER"
    peach = dominant(p4, (430, 175, 500, 238), (255, 255, 255), 20)  # big outlined numeral
    muted = dominant(p8, (48, 30, 133, 39), (255, 255, 255), 30)  # breadcrumb grey (white header band)
    palette = {
        'dk1': dark, 'lt1': (255, 255, 255), 'dk2': grey, 'lt2': bg,
        'accent1': red, 'accent2': orange, 'accent3': grey, 'accent4': pink,
        'accent5': salmon, 'accent6': peach, 'hlink': red, 'folHlink': muted,
    }
    print('palette:', {k: '#' + hexof(v) for k, v in palette.items()})

    # ---- crops (cover art without the title/date, logo from cover bottom-left)
    art_box = (376, 0, 960, 486)
    art = keyed_crop(p1, art_box, bg, blank=[(376, 355, 600, 430)])
    art_png = os.path.join(out_dir, 'halftone-cover.png')
    art.save(art_png)
    logo_box = (47, 468, 191, 503)
    logo = keyed_crop(p1, logo_box, bg)
    logo_png = os.path.join(out_dir, 'logo.png')
    logo.save(logo_png)

    # ---- presentation
    prs = Presentation()
    old_w = prs.slide_width
    prs.slide_width, prs.slide_height = 12192000, 6858000
    f = prs.slide_width / old_w
    master = prs.slide_master
    scale_x(master._element, f)
    for l in master.slide_layouts:
        scale_x(l._element, f)

    # theme: colours + fonts
    theme_part = master.part.part_related_by(RT.THEME)
    theme = etree.fromstring(theme_part.blob)
    ns = {'a': A}
    theme.find('.//a:clrScheme', ns).set('name', 'SK hynix (reconstructed)')
    for k, rgb in palette.items():
        slot = theme.find(f'.//a:clrScheme/a:{k}', ns)
        for c in list(slot):
            slot.remove(c)
        etree.SubElement(slot, f'{{{A}}}srgbClr', val=hexof(rgb))
    fs = theme.find('.//a:fontScheme', ns)
    fs.set('name', 'SK hynix (reconstructed)')
    for tag, face in (('majorFont', FONT_MAJOR), ('minorFont', FONT_MINOR)):
        el = fs.find(f'a:{tag}', ns)
        el.find('a:latin', ns).set('typeface', face)
        el.find('a:ea', ns).set('typeface', face)
        for extra in el.findall('a:font', ns):
            if extra.get('script') == 'Hang':
                extra.set('typeface', face)
    theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)

    # master: off-white background, text defaults
    set_bg(master._element, 'bg2')  # lt2 = sampled off-white
    tx = master._element.find(qn('p:txStyles'))
    t1 = tx.find(qn('p:titleStyle')).find(qn('a:lvl1pPr')).find(qn('a:defRPr'))
    t1.set('sz', '2200'); t1.set('b', '1')
    b1 = tx.find(qn('p:bodyStyle')).find(qn('a:lvl1pPr')).find(qn('a:defRPr'))
    b1.set('sz', '1600')

    # keep five layouts, drop the rest
    keep = {'Title Slide', 'Title and Content', 'Section Header', 'Title Only', 'Blank'}
    for l in list(master.slide_layouts):
        if l.name not in keep:
            master.slide_layouts.remove(l)
    L = {l.name: l for l in master.slide_layouts}

    # Title Slide (cover, PDF page 1): halftone art right, logo bottom-left, title/date bottom-left
    cov = L['Title Slide']
    drop_placeholders(cov, {'ctrTitle', 'subTitle'})
    t = placeholder(cov, 'ctrTitle')
    set_xfrm(t, 48, 362, 600, 40)
    set_lst_style(t, size=28, bold=True, scheme='accent1', algn='l')
    st = placeholder(cov, 'subTitle')
    set_xfrm(st, 52, 402, 400, 26)
    set_lst_style(st, size=14, bold=False, scheme='tx1', algn='l')
    add_picture(cov, art_png, *art_box[:2], art_box[2] - art_box[0], art_box[3] - art_box[1], 'Halftone art')
    add_picture(cov, logo_png, logo_box[0], logo_box[1], logo_box[2] - logo_box[0], logo_box[3] - logo_box[1], 'Logo')

    # Section Header (chapter divider, PDF page 4): white, centred title, orange rule, bullets below
    sec = L['Section Header']
    set_bg(sec._element, 'bg1')
    drop_placeholders(sec, {'title', 'body'})
    t = placeholder(sec, 'title')
    set_xfrm(t, 180, 252, 600, 66)
    set_lst_style(t, size=48, bold=True, scheme='accent1', algn='ctr')
    b = placeholder(sec, 'body')
    set_xfrm(b, 395, 322, 330, 120)
    set_lst_style(b, size=14, bold=False, scheme='tx1', algn='l')
    add_rect(sec, 302, 239.5, 355, 1, 'accent2', 'Chapter rule')

    # Title and Content / Title Only (PDF pages 5-12): white header band, title top-left
    for name in ('Title and Content', 'Title Only'):
        lay = L[name]
        drop_placeholders(lay, {'title', 'body', 'obj'})
        t = placeholder(lay, 'title')
        set_xfrm(t, 48, 42, 864, 34)
        set_lst_style(t, size=22, bold=True, scheme='tx1', algn='l')
        body = placeholder(lay, 'body', 'obj')
        if body is not None:
            set_xfrm(body, 48, 100, 864, 380)
            set_lst_style(body, size=16, bold=False, scheme='tx1')
        add_rect(lay, 0, 0, 960, 83, 'bg1', 'Header band')

    # logo bottom-right on every non-cover layout (PDF pages 2-12)
    for name in ('Section Header', 'Title and Content', 'Title Only', 'Blank'):
        lay = L[name]
        if name == 'Blank':
            drop_placeholders(lay, set())
        # same crop, at the smaller size of the bottom-right logo image in the PDF (794,489 118x27)
        w = 118
        add_picture(lay, logo_png, 794, 488, w, w * (logo_box[3] - logo_box[1]) / (logo_box[2] - logo_box[0]), 'Logo')

    # order: cover, section, content, title only, blank
    lst = master._element.find(qn('p:sldLayoutIdLst'))
    order = ['Title Slide', 'Section Header', 'Title and Content', 'Title Only', 'Blank']
    ids = {master.part.related_part(e.get(qn('r:id'))).slide_layout.name: e for e in lst}
    for e in list(lst):
        lst.remove(e)
    for n in order:
        lst.append(ids[n])

    out = os.path.join(out_dir, 'skhynix-reconstructed.pptx')
    prs.save(out)
    print('wrote', out)


if __name__ == '__main__':
    main()
